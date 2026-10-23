"""
Synthèse des familles de mutation : test sur les marches V1 (teste/parametre/{Famille}.csv) et énumération
exhaustive (teste/parametre/exhaustif/{Famille}.csv) dans un seul tableau, deux lignes par sous-famille (une par
source), avec intervalles de confiance, règles qui suffisent seules comme filtre, statut (accord / désaccord / une
seule source) et lien vers l'image de la sous-famille.

Écrit teste/parametre/synthese_familles.csv, synthese_familles.xlsx (liens cliquables, couleurs, onglets « clé de
lecture » et « méthode » ; si le .xlsx existe déjà, synthese_familles_nouveau.xlsx, pour ne pas écraser les réglages
faits à la main) et synthese_familles_methode.md. Relancer après avoir refait des tests :
    /Applications/anaconda3/envs/revomol/bin/python teste/synthese_familles.py

Le .xlsx est écrit par le Python d'Anaconda de base (openpyxl), l'environnement revomol n'ayant pas de module Excel.
"""

from __future__ import annotations

import math
import os
import re
import subprocess
import sys
from typing import Any

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DIR = os.path.join(HERE, "parametre")
FAMILIES = ["AddAtomMG", "ChangeBondMG", "SubstituteShiftMG", "RemoveAtomMG", "CutAtomMG", "InsertCarbonMG"]
RULES = ["R1", "R2a", "R2b", "R2c", "R3", "R4", "R5"]
JEUX = {
    "marches": "marches V1 (test_Regle_V1), 10 000 molécules conformes tirées (graine 0)",
    "exhaustif": "énumération de 1 à 8 atomes, 460 533 molécules (8 atomes : ancienne R2, O seul)",
}
PYTHON_EXCEL = "/Applications/anaconda3/bin/python"
Z = 1.959963984540054  # quantile 97,5 % de la loi normale (intervalle à 95 %)

ALWAYS, NEVER, DEPENDS = "toujours autorisée", "toujours interdite", "dépend"


# ------------------------------------------------------------------------------ formules
def wilson(k: int, n: int) -> tuple[float, float]:
    """Intervalle de confiance à 95 % (Wilson) d'une proportion k / n, en %."""
    if n == 0:
        return math.nan, math.nan
    p = k / n
    centre = (p + Z * Z / (2 * n)) / (1 + Z * Z / n)
    demi = Z * math.sqrt(p * (1 - p) / n + Z * Z / (4 * n * n)) / (1 + Z * Z / n)
    return round(100 * max(0.0, centre - demi), 2), round(100 * min(1.0, centre + demi), 2)


def zero_bound(n: int) -> float:
    """Avec 0 exception sur n essais : taux d'exception maximal plausible à 95 %, en % (1 − 0,05^(1/n) ≈ 3/n)."""
    return round(100 * (1 - 0.05 ** (1 / n)), 4) if n else math.nan


def verdict(accepted: int, total: int) -> str:
    return ALWAYS if accepted == total else NEVER if accepted == 0 else DEPENDS


# ------------------------------------------------------------------------------ lecture
def parse_counts(text: Any) -> dict[str, int]:
    """« R2a ×5, R3 » → {"R2a": 5, "R3": 1} ; « — » ou vide → {}."""
    out: dict[str, int] = {}
    if not isinstance(text, str) or text.strip() in ("", "—"):
        return out
    for part in text.split(","):
        m = re.fullmatch(r"\s*(\S+?)(?:\s*×\s*(\d+))?\s*", part)
        if m:
            out[m.group(1)] = int(m.group(2) or 1)
    return out


def subfamily_file_name(name: str) -> str:
    """Nom de fichier d'une image (comme parametrage_regles.subfamily_file_name)."""
    return re.sub(r"[^\w@.+-]+", "_", name.replace("→", "-"))


def read_source(source: str, family: str) -> pd.DataFrame:
    path = os.path.join(DIR, family + ".csv") if source == "marches" else os.path.join(DIR, "exhaustif", family + ".csv")
    if not os.path.exists(path):
        return pd.DataFrame()
    table = pd.read_csv(path)
    rows = []
    for r in table.to_dict("records"):
        n, acc, ref = int(r["candidates"]), int(r["acceptées"]), int(r["refusées"])
        per_rule = ({rule: int(r[rule]) for rule in RULES if rule in r} if source == "marches"
                    else {rule: parse_counts(r.get("règles qui refusent")).get(rule, 0) for rule in RULES})
        alone = parse_counts(r.get("règle seule (refus)"))
        rows.append({"famille": r["famille"], "sous-famille": r["sous-famille"], "source": source, "n": n, "acc": acc,
                     "ref": ref, "per_rule": per_rule, "alone": alone, "molécules": r.get("molécules"),
                     "min_acc": r.get("plus petite acceptée (atomes)"), "min_ref": r.get("plus petite refusée (atomes)")})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------ une ligne
def line(r: dict) -> dict:
    n, acc, ref = r["n"], r["acc"], r["ref"]
    low, high = wilson(acc, n)
    v = verdict(acc, n)
    out: dict[str, Any] = {
        "famille": r["famille"], "sous-famille": r["sous-famille"], "source": r["source"], "jeu": JEUX[r["source"]],
        "candidates": n, "molécules": r["molécules"], "acceptées": acc, "refusées": ref,
        "% acceptées": round(100 * acc / n, 2) if n else math.nan, "IC 95 % bas (%)": low, "IC 95 % haut (%)": high,
        "verdict": v, "exception max 95 % (%)": zero_bound(n) if v != DEPENDS else math.nan,
    }
    for rule in RULES:
        out[f"refus {rule}"] = r["per_rule"].get(rule, 0)
    out["refus par une seule règle"] = sum(r["alone"].values())
    for rule in RULES:
        out[f"seule {rule}"] = r["alone"].get(rule, 0)
    # une règle suffit seule comme filtre si elle est cassée par tous les refus
    if ref:
        coverage = {rule: r["per_rule"].get(rule, 0) for rule in RULES}
        sufficient = [rule for rule in RULES if coverage[rule] == ref]
        best = max(RULES, key=lambda rule: coverage[rule])
        b_low, b_high = wilson(coverage[best], ref)
        out["règle(s) qui suffisent seules"] = ", ".join(sufficient) or "aucune"
        out["oubli max 95 % (%)"] = zero_bound(ref) if sufficient else math.nan
        out["meilleure règle seule"] = best
        out["couverture (%)"] = round(100 * coverage[best] / ref, 2)
        out["couverture IC 95 % bas (%)"] = b_low
        out["couverture IC 95 % haut (%)"] = b_high
        out["règles indispensables"] = ", ".join(rule for rule in RULES if r["alone"].get(rule)) or "aucune"
    else:
        out.update({"règle(s) qui suffisent seules": "—", "oubli max 95 % (%)": math.nan, "meilleure règle seule": "—",
                    "couverture (%)": math.nan, "couverture IC 95 % bas (%)": math.nan,
                    "couverture IC 95 % haut (%)": math.nan, "règles indispensables": "—"})
    out["plus petite acceptée (atomes)"] = r["min_acc"]
    out["plus petite refusée (atomes)"] = r["min_ref"]
    return out


def build() -> tuple[pd.DataFrame, dict]:
    lines, missing_images = [], 0
    stats = {"familles": [], "sous-familles": 0}
    for family in FAMILIES:
        parts = {s: read_source(s, family) for s in ("marches", "exhaustif")}
        if all(p.empty for p in parts.values()):
            continue
        stats["familles"].append(family)
        by_source = {s: {r["sous-famille"]: r for r in p.to_dict("records")} if not p.empty else {}
                     for s, p in parts.items()}
        names = sorted(set(by_source["marches"]) | set(by_source["exhaustif"]))
        for name in names:
            stats["sous-familles"] += 1
            here = {s: line(by_source[s][name]) for s in ("marches", "exhaustif") if name in by_source[s]}
            if len(here) == 1:
                status = "seulement " + next(iter(here))
                significant = ""
            else:
                m, e = here["marches"], here["exhaustif"]
                status = "accord" if m["verdict"] == e["verdict"] else "désaccord"
                disjoint = m["IC 95 % haut (%)"] < e["IC 95 % bas (%)"] or e["IC 95 % haut (%)"] < m["IC 95 % bas (%)"]
                significant = "oui" if disjoint else "non"
            for source, out in here.items():
                out["statut"] = status
                out["écart significatif (IC 95 % disjoints)"] = significant
                rel = (os.path.join(family, subfamily_file_name(name) + ".png") if source == "marches"
                       else os.path.join("exhaustif", family, subfamily_file_name(name) + ".png"))
                if os.path.exists(os.path.join(DIR, rel)):
                    out["image"] = rel
                else:
                    out["image"] = ""
                    missing_images += 1
                lines.append(out)
    table = pd.DataFrame(lines)
    stats["images manquantes"] = missing_images
    return table, stats


METHODE = """# Synthèse des familles de mutation : clé de lecture et méthode

## Clé de lecture
Une ligne = une sous-famille de mutation vue par une source : « marches » (10 000 molécules des marches V1) ou
« exhaustif » (toutes les molécules de 1 à 8 atomes). Les deux lignes d'une sous-famille se suivent. Lire une ligne
dans cet ordre :
1. **verdict** : toujours autorisée (aucune mutation refusée), toujours interdite (toutes refusées) ou dépend.
2. **est-ce sûr ?** « exception max 95 % » (pour un « toujours ») ou la largeur de « IC 95 % » (pour un « dépend ») :
   dans le .xlsx, rouge = sûr, bleu = trop peu de cas pour conclure (pourrait être le hasard).
3. **qui refuse ?** « refus Rx » = refus où la règle Rx est cassée ; « seule Rx » = refus où Rx est la seule cassée
   (sans Rx, ces mutations passeraient).
4. **une règle suffit-elle seule ?** « règle(s) qui suffisent seules » : cassée par 100 % des refus, n'appliquer
   qu'elle donnerait le même résultat ; « oubli max 95 % » dit à quel point c'est sûr ; « aucune » : il en faut
   plusieurs (voir « meilleure règle seule » et sa « couverture »).
5. **les deux sources sont-elles d'accord ?** « statut ».

Exemples :
- addCs@Cs : toujours interdite partout ; exception max 0,027 % (marches, 11 065 cas), 0,0008 % (exhaustif) → sûr.
  R2a est cassée par 100 % des refus → R2a suffit seule ; R5 est aussi cassée (5 034 cas) mais jamais seule.
- subOs→Nd+mont : dépend, 65,7 % acceptées, IC 64,8–66,6 % → proportion bien connue ; tous les refus cassent R3 →
  R3 suffit seule (oubli max 0,08 %).
- cutCs@Cs,Os : marches 80,95 % sur 21 cas (IC 60,0–92,3 %, peu sûr) ; exhaustif toujours autorisée (0 refus sur
  2 027). Désaccord : les 4 refus des marches viennent de R4 (grands cycles absents à 8 atomes) ; oubli max 52,7 % :
  4 refus ne suffisent pas pour dire que R4 suffirait toujours.
- cutCt@Ct,Ct (marches) : 2 candidates → « toujours autorisée » mais exception max 77,6 % : aucune conclusion.

Couleurs du .xlsx : verdict et statut (vert = autorisée / accord, rouge = interdite / désaccord, jaune = dépend / une
seule source) ; colonnes des règles dans la couleur de la règle, pâle si elle ne refuse rien, plus foncé si elle
refuse ; statistiques de certitude de rouge (≤ 0,01 %, sûr) à bleu (≥ 50 %, le hasard), échelle logarithmique.

# Méthode

Fichiers : `synthese_familles.csv` et `synthese_familles.xlsx` (même tableau, liens cliquables dans le .xlsx),
écrits par `teste/synthese_familles.py` à partir des CSV de `teste/parametre/` (marches) et
`teste/parametre/exhaustif/` (énumération).

## Les deux sources
- **marches** : 10 000 molécules conformes tirées dans les marches `test_Regle_V1` (graine 0) ; toutes les
  mutations candidates de la famille sont jugées par les règles du notebook.
- **exhaustif** : toutes les molécules C, N, O conformes de 1 à 8 atomes (460 533 ; celles de 8 atomes viennent de
  l'énumération avec l'ancienne R2, O seul dans l'exception du cycle).
Deux lignes par sous-famille, une par source ; une sous-famille absente d'une source n'a qu'une ligne.

## Intervalle de confiance d'un % (colonnes « IC 95 % »)
Pour k mutations acceptées sur n candidates, p = k / n, l'intervalle de Wilson à 95 % :

    centre = (p + z²/(2n)) / (1 + z²/n)
    demi-largeur = z · √( p(1−p)/n + z²/(4n²) ) / (1 + z²/n),   z = 1,96

On n'utilise pas p ± 1,96·√(p(1−p)/n) : cette formule donne un intervalle nul quand p vaut 0 ou 100 %, et des
bornes hors de [0, 100 %] près des extrêmes.

## Verdict « toujours » (colonne « exception max 95 % »)
Si aucune des n candidates ne fait exception (0 % ou 100 % acceptées), le vrai taux d'exception est, à 95 %, au
plus 1 − 0,05^(1/n), soit environ 3/n (« règle de trois »). Ex. : 0 exception sur 1 000 → au plus 0,30 %.

## Quelle règle refuse quoi
- « refus Rx » : nombre de refus où la règle Rx est cassée (une mutation peut en casser plusieurs : la somme peut
  dépasser « refusées »).
- « refus par une seule règle » et « seule Rx » : refus où Rx est la **seule** règle cassée. Une règle avec
  « seule Rx » > 0 est **indispensable** pour cette sous-famille : sans elle, ces mutations passeraient.

## Une règle suffit-elle seule comme filtre ?
Couverture de Rx = refus qui cassent Rx / refusées.
- Couverture 100 % : appliquer **seulement Rx** à cette sous-famille refuserait exactement les mêmes mutations
  (« règle(s) qui suffisent seules »). Avec 0 oubli sur `refusées` cas, le taux d'oubli réel est, à 95 %, au plus
  1 − 0,05^(1/refusées) (« oubli max 95 % »).
- Sinon : la meilleure règle seule, sa couverture et son intervalle de Wilson ; les refus qu'elle ne couvre pas
  passeraient si on n'appliquait qu'elle.

## Statut et écart
- **accord** : même verdict dans les deux sources ; **désaccord** : verdicts différents ;
  **seulement exhaustif / seulement marches** : la sous-famille n'existe que dans une source.
- **écart significatif** : les deux intervalles à 95 % du % acceptées ne se recouvrent pas.

## Limites
- L'**exhaustif n'est pas un échantillon** : jusqu'à 8 atomes, ses chiffres sont exacts. Son incertitude porte sur
  les molécules plus grandes, qu'il ne contient pas (voir la stabilisation dans `contextes_mutations.ipynb`).
- Les mutations d'une même molécule **ne sont pas indépendantes** : les intervalles sont un peu trop étroits.
- Les **marches V1** sont un échantillon biaisé (5 types de mutation, départ benzène) : un écart avec l'exhaustif
  vient souvent de contextes absents des marches ou de molécules plus grandes que 8 atomes.
- Les deux sources n'ont pas la même population : un écart de % n'est pas une erreur, c'est surtout le verdict
  (accord / désaccord) qui se compare.
"""


EXCEL_SCRIPT = r'''
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

csv_path, xlsx_path, md_path = sys.argv[1:4]
table = pd.read_csv(csv_path, keep_default_na=False)
wb = Workbook()
ws = wb.active
ws.title = "synthèse"
ws.append(list(table.columns))
for row in table.itertuples(index=False):
    ws.append([None if v == "" else v for v in row])
colours = {"accord": "D8EFD3", "désaccord": "F8D0D0", "seulement exhaustif": "FFF2B3", "seulement marches": "D6E4F5"}
cols = {name: i + 1 for i, name in enumerate(table.columns)}
for r in range(2, ws.max_row + 1):
    status = ws.cell(r, cols["statut"]).value
    if status in colours:
        ws.cell(r, cols["statut"]).fill = PatternFill("solid", fgColor=colours[status])
    v = ws.cell(r, cols["verdict"]).value
    ws.cell(r, cols["verdict"]).fill = PatternFill("solid", fgColor={"toujours autorisée": "D8EFD3",
        "toujours interdite": "F8D0D0", "dépend": "FFF2B3"}.get(v, "FFFFFF"))
    cell = ws.cell(r, cols["image"])
    if cell.value:
        cell.hyperlink = cell.value
        cell.value = "voir l'image"
        cell.font = Font(color="0563C1", underline="single")
for c in range(1, ws.max_column + 1):
    ws.cell(1, c).font = Font(bold=True)
    ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[get_column_letter(c)].width = 14
for name, width in (("famille", 20), ("sous-famille", 18), ("jeu", 30), ("règle(s) qui suffisent seules", 18),
                    ("règles indispensables", 18), ("statut", 18)):
    ws.column_dimensions[get_column_letter(cols[name])].width = width
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions
method = wb.create_sheet("méthode")
method.column_dimensions["A"].width = 120
for text in open(md_path, encoding="utf-8").read().split("\n"):
    method.append([text])
    if text.startswith("#"):
        method.cell(method.max_row, 1).font = Font(bold=True)
wb.save(xlsx_path)
'''


def main() -> None:
    table, stats = build()
    csv_path = os.path.join(DIR, "synthese_familles.csv")
    xlsx_path = os.path.join(DIR, "synthese_familles.xlsx")
    if os.path.exists(xlsx_path):
        # ne pas écraser les réglages faits à la main dans le .xlsx existant
        xlsx_path = os.path.join(DIR, "synthese_familles_nouveau.xlsx")
    md_path = os.path.join(DIR, "synthese_familles_methode.md")
    table.to_csv(csv_path, index=False)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(METHODE)
    result = subprocess.run([PYTHON_EXCEL, "-c", EXCEL_SCRIPT, csv_path, xlsx_path, md_path], capture_output=True,
                            text=True, check=False)
    if result.returncode != 0:
        print("le .xlsx n'a pas pu être écrit :", result.stderr[-800:], file=sys.stderr)
    else:  # couleurs et clé de lecture
        subprocess.run([PYTHON_EXCEL, os.path.join(HERE, "colorer_synthese.py"), xlsx_path], check=False)
    counts = table.drop_duplicates(["famille", "sous-famille"])["statut"].value_counts().to_dict()
    print(f"{len(table)} lignes ({stats['sous-familles']} sous-familles, familles : {', '.join(stats['familles'])})")
    print("statuts :", counts, "· images manquantes :", stats["images manquantes"])
    print("écrit :", csv_path, xlsx_path if result.returncode == 0 else "(pas de .xlsx)", md_path, sep="\n  ")


if __name__ == "__main__":
    main()
